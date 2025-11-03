"""
Servicio de DJ Integral para Evolve Soluciones
==============================================

Contiene toda la lógica de negocio para la funcionalidad de Declaraciones Juradas (DJ),
incluyendo obtención de datos, filtros y exportación.
"""

import pymysql
from datetime import datetime
import re
from typing import Dict, List, Optional, Any

from aplicacion.modelos.base_datos import obtener_conexion_local
from configuracion.configuracion import obtener_configuracion

# Obtener configuración
config = obtener_configuracion()


class ServicioDJIntegral:
    """Servicio para manejar consultas de Declaraciones Juradas Integrales"""

    def __init__(self, base_datos='stratex'):
        """
        Inicializa el servicio con la base de datos a usar

        Args:
            base_datos (str): Nombre de la base de datos MySQL (ej: 'stratex', 'evolve', etc.)
        """
        self.config = config
        self.base_datos = base_datos

    def obtener_conexion(self):
        """
        Obtiene conexión a la base de datos MySQL asignada al usuario

        Returns:
            pymysql.Connection: Conexión a la base de datos
        """
        return obtener_conexion_local(self.base_datos)

    def limpiar_rut(self, rut):
        """
        Limpia el RUT removiendo puntos y guiones

        Args:
            rut (str): RUT con formato

        Returns:
            str: RUT limpio sin puntos ni guiones
        """
        if not rut:
            return ""
        return re.sub(r'[.\-]', '', str(rut))

    def obtener_datos_empresas_con_dj(self, filtros=None):
        """
        Obtiene datos de empresas con sus Declaraciones Juradas por año

        Args:
            filtros (dict): Filtros a aplicar en la consulta

        Returns:
            dict: Diccionario con empresas agrupadas y años disponibles
        """
        if filtros is None:
            filtros = {}

        conexion = None
        try:
            conexion = self.obtener_conexion()

            # Construir consulta base
            consulta_base = f"""
                SELECT DISTINCT
                    e.run_rut as rut,
                    e.empresa as nombre,
                    e.auditor as usuario,
                    e.grupo as grupo
                FROM empresas e
                WHERE 1=1
            """

            parametros = []

            # Aplicar filtros de empresa
            if filtros.get('empresa_filtro'):
                consulta_base += " AND e.empresa LIKE %s"
                parametros.append(f"%{filtros['empresa_filtro']}%")

            if filtros.get('rut_filtro'):
                rut_limpio = self.limpiar_rut(filtros['rut_filtro'])
                consulta_base += " AND (e.run_rut LIKE %s OR e.run_rut LIKE %s)"
                parametros.extend([f"%{filtros['rut_filtro']}%", f"%{rut_limpio}%"])

            if filtros.get('usuario_filtro'):
                consulta_base += " AND e.auditor = %s"
                parametros.append(filtros['usuario_filtro'])

            if filtros.get('grupo_filtro'):
                consulta_base += " AND e.grupo = %s"
                parametros.append(filtros['grupo_filtro'])

            consulta_base += " ORDER BY e.empresa"

            with conexion.cursor(pymysql.cursors.DictCursor) as cursor:
                cursor.execute(consulta_base, parametros)
                empresas = cursor.fetchall()

            # Subconsulta para obtener la última fecha de consulta por RUT
            ruts_empresas = [e['rut'] for e in empresas]
            fechas_consulta = {}
            if ruts_empresas:
                # Usamos un placeholder %s para cada RUT para seguridad
                placeholders = ', '.join(['%s'] * len(ruts_empresas))
                consulta_fechas = f"""
                    SELECT rut, MAX(DATE(fecha_consulta)) as ultima_fecha
                    FROM dj_integral
                    WHERE rut IN ({placeholders}) AND estado = 'T'
                    GROUP BY rut
                """
                print(f"DEBUG DJ FECHAS: Consultando fechas para RUTs: {ruts_empresas}")
                print(f"DEBUG DJ FECHAS: Query: {consulta_fechas}")
                with conexion.cursor(pymysql.cursors.DictCursor) as cursor:
                    cursor.execute(consulta_fechas, ruts_empresas)
                    resultados_fechas = cursor.fetchall()
                    print(f"DEBUG DJ FECHAS: Resultados: {resultados_fechas}")
                    for fila in resultados_fechas:
                        fechas_consulta[fila['rut']] = fila['ultima_fecha']
                        print(f"DEBUG DJ FECHAS: RUT {fila['rut']} -> Fecha {fila['ultima_fecha']}")
                print(f"DEBUG DJ FECHAS: Dict final fechas_consulta: {fechas_consulta}")

            # Agrupar resultados por empresa con DJ por año
            empresas_agrupadas = {}
            años_disponibles = set()

            for empresa in empresas:
                rut = empresa['rut']

                # Obtener DJ de esta empresa
                dj_empresa = self.obtener_dj_por_rut(rut)

                print(f"DEBUG: Empresa {empresa['nombre']} (RUT: {rut}) - DJ encontrados: {len(dj_empresa)}")

                empresas_agrupadas[rut] = {
                    'rut': rut,
                    'nombre': empresa['nombre'],
                    'usuario': empresa['usuario'],
                    'grupo': empresa['grupo'],
                    'ultima_fecha_consulta': fechas_consulta.get(rut),  # Añadir fecha aquí
                    'dj_por_anio': {},
                    'total_observadas': 0  # Inicializar contador
                }

                # Organizar DJ por año
                for dj in dj_empresa:
                    # Determinar qué años tienen datos
                    for año in ['2025', '2024', '2023', '2022', '2021', '2020']:
                        estado_campo = f'estado_{año}'
                        estado = dj.get(estado_campo)

                        print(f"DEBUG: DJ {dj['dj_numero']} - Año {año} - Campo {estado_campo} - Estado: {estado}")

                        if estado:  # Si tiene estado en este año
                            años_disponibles.add(int(año))

                            if año not in empresas_agrupadas[rut]['dj_por_anio']:
                                empresas_agrupadas[rut]['dj_por_anio'][año] = []

                            # Contar observadas
                            if estado == 'Observada':
                                empresas_agrupadas[rut]['total_observadas'] += 1

                            # Verificar si este DJ ya existe en este año
                            dj_existente = False
                            for dj_guardado in empresas_agrupadas[rut]['dj_por_anio'][año]:
                                if dj_guardado['dj_numero'] == dj['dj_numero']:
                                    dj_existente = True
                                    # Actualizar el estado si es diferente
                                    if dj_guardado['estado'] != estado:
                                        dj_guardado['estado'] = estado
                                    break

                            if not dj_existente:
                                empresas_agrupadas[rut]['dj_por_anio'][año].append({
                                    'dj_numero': dj['dj_numero'],
                                    'titulo': dj['titulo'],
                                    'estado': estado,
                                    'dj_id': dj['dj_id']
                                })
                                print(f"DEBUG: Agregado DJ {dj['dj_numero']} al año {año} con estado {estado}")

            # Convertir años_disponibles a lista ordenada descendente
            años_ordenados = sorted(list(años_disponibles), reverse=True)

            return {
                'empresas': empresas_agrupadas,
                'años': años_ordenados
            }

        except Exception as e:
            print(f"ERROR: Error obteniendo datos de empresas con DJ: {e}")
            import traceback
            traceback.print_exc()
            return {
                'empresas': {},
                'años': []
            }
        finally:
            if conexion:
                conexion.close()

    def obtener_dj_por_rut(self, rut: str) -> List[Dict[str, Any]]:
        """
        Obtiene todas las Declaraciones Juradas vigentes de un RUT específico

        Args:
            rut (str): RUT de la empresa

        Returns:
            List[Dict]: Lista de DJ con sus estados por año
        """
        conexion = None
        try:
            conexion = self.obtener_conexion()

            consulta = f"""
                SELECT
                    dj_numero,
                    titulo,
                    dj_id,
                    estado_2025,
                    estado_2024,
                    estado_2023,
                    estado_2022,
                    estado_2021,
                    estado_2020
                FROM {self.base_datos}.dj_integral
                WHERE rut = %s AND estado = 'T'
                ORDER BY dj_numero
            """

            with conexion.cursor(pymysql.cursors.DictCursor) as cursor:
                cursor.execute(consulta, [rut])
                resultados = cursor.fetchall()

                print(f"DEBUG: Buscando DJ para RUT: {rut}")
                print(f"DEBUG: Query: {consulta}")
                print(f"DEBUG: Resultados encontrados: {len(resultados) if resultados else 0}")
                if resultados:
                    print(f"DEBUG: Primer resultado: {resultados[0]}")

                return list(resultados) if resultados else []

        except Exception as e:
            print(f"ERROR: Error obteniendo DJ para RUT {rut}: {e}")
            import traceback
            traceback.print_exc()
            return []
        finally:
            if conexion:
                conexion.close()

    def obtener_usuarios_unicos(self) -> List[str]:
        """
        Obtiene lista de usuarios únicos (auditores) para filtros

        Returns:
            List[str]: Lista de auditores únicos
        """
        conexion = None
        try:
            conexion = self.obtener_conexion()

            with conexion.cursor() as cursor:
                cursor.execute("""
                    SELECT DISTINCT auditor
                    FROM empresas
                    WHERE auditor IS NOT NULL
                    ORDER BY auditor
                """)
                resultados = cursor.fetchall()

                return [fila[0] for fila in resultados if fila[0]]

        except Exception as e:
            print(f"ERROR: Error obteniendo usuarios únicos: {e}")
            return []
        finally:
            if conexion:
                conexion.close()

    def obtener_grupos_unicos(self) -> List[str]:
        """
        Obtiene lista de grupos únicos para filtros

        Returns:
            List[str]: Lista de grupos únicos
        """
        conexion = None
        try:
            conexion = self.obtener_conexion()

            with conexion.cursor() as cursor:
                cursor.execute("""
                    SELECT DISTINCT grupo
                    FROM empresas
                    WHERE grupo IS NOT NULL
                    ORDER BY grupo
                """)
                resultados = cursor.fetchall()

                return [fila[0] for fila in resultados if fila[0]]

        except Exception as e:
            print(f"ERROR: Error obteniendo grupos únicos: {e}")
            return []
        finally:
            if conexion:
                conexion.close()

    def obtener_estados_unicos(self) -> List[str]:
        """
        Obtiene lista de estados únicos presentes en DJ

        Returns:
            List[str]: Lista de estados únicos
        """
        conexion = None
        try:
            conexion = self.obtener_conexion()

            # Obtener todos los estados de todos los años
            estados = set()

            consulta = f"""
                SELECT DISTINCT
                    estado_2025, estado_2024, estado_2023,
                    estado_2022, estado_2021, estado_2020
                FROM {self.base_datos}.dj_integral
                WHERE estado = 'T'
            """

            with conexion.cursor(pymysql.cursors.DictCursor) as cursor:
                cursor.execute(consulta)
                resultados = cursor.fetchall()

                for fila in resultados:
                    for campo, valor in fila.items():
                        if valor:
                            estados.add(valor)

            return sorted(list(estados))

        except Exception as e:
            print(f"ERROR: Error obteniendo estados únicos: {e}")
            return []
        finally:
            if conexion:
                conexion.close()

    def obtener_estadisticas(self) -> Dict[str, Any]:
        """
        Obtiene estadísticas generales de DJ

        Returns:
            Dict: Estadísticas de DJ por empresa
        """
        conexion = None
        try:
            conexion = self.obtener_conexion()

            with conexion.cursor(pymysql.cursors.DictCursor) as cursor:
                # Total de empresas con DJ
                cursor.execute(f"""
                    SELECT COUNT(DISTINCT rut) as total
                    FROM {self.base_datos}.dj_integral
                    WHERE estado = 'T'
                """)
                resultado_total_empresas = cursor.fetchone()
                total_empresas_con_dj = resultado_total_empresas['total'] if resultado_total_empresas else 0

                # Total de DJ vigentes
                cursor.execute(f"""
                    SELECT COUNT(*) as total
                    FROM {self.base_datos}.dj_integral
                    WHERE estado = 'T'
                """)
                resultado_total_dj = cursor.fetchone()
                total_dj = resultado_total_dj['total'] if resultado_total_dj else 0

                # DJ por número
                cursor.execute(f"""
                    SELECT dj_numero, COUNT(*) as cantidad
                    FROM {self.base_datos}.dj_integral
                    WHERE estado = 'T'
                    GROUP BY dj_numero
                    ORDER BY dj_numero
                """)
                por_numero = cursor.fetchall()

                return {
                    'total_empresas_con_dj': total_empresas_con_dj,
                    'total_dj': total_dj,
                    'por_numero': por_numero
                }

        except Exception as e:
            print(f"ERROR: Error obteniendo estadísticas: {e}")
            return {
                'total_empresas_con_dj': 0,
                'total_dj': 0,
                'por_numero': []
            }
        finally:
            if conexion:
                conexion.close()

    def obtener_observaciones_dj(self, rut: str, dj_numero: int, periodo: int) -> List[Dict[str, Any]]:
        """
        Obtiene observaciones asociadas a una declaración jurada (DJ) específica.

        Args:
            rut (str): RUT de la empresa
            dj_numero (int): Número de la DJ
            periodo (int): Año del período (ej: 2022)

        Returns:
            List[Dict]: Lista de observaciones (code, descripcion, orientacion)
        """
        conexion = None
        try:
            conexion = self.obtener_conexion()
            consulta = f"""
                SELECT a.observacion_code, a.glosa_observacion, a.descripcion, a.orientacion
                FROM {self.base_datos}.dj_integral_observaciones a
                WHERE a.rut = %s AND a.dj_numero = %s AND a.periodo = %s
                ORDER BY a.observacion_code
            """

            with conexion.cursor(pymysql.cursors.DictCursor) as cursor:
                cursor.execute(consulta, (rut, dj_numero, periodo))
                resultados = cursor.fetchall()
                return list(resultados) if resultados else []

        except Exception as e:
            print(f"ERROR: Error obteniendo observaciones DJ [{rut} - {dj_numero} - {periodo}]: {e}")
            import traceback
            traceback.print_exc()
            return []
        finally:
            if conexion:
                conexion.close()

    def exportar_observaciones_dj_excel(self, nombre_usuario: str, es_administrador: bool = False):
        """
        Exporta observaciones relacionadas con DJ a un archivo Excel en memoria.

        Args:
            nombre_usuario (str): Nombre del usuario que solicita la exportación
            es_administrador (bool): Si True, exporta todas las observaciones; si False, solo las del usuario

        Returns:
            BytesIO: Buffer con el archivo Excel listo para enviar como attachment, o None si no hay datos
        """
        try:
            # Import tardío para no forzar dependencia si no se usa
            from io import BytesIO
            from openpyxl import Workbook

            conexion = self.obtener_conexion()
            with conexion.cursor(pymysql.cursors.DictCursor) as cursor:
                # Si no es admin, filtrar por auditor/usuario en la tabla empresas
                if not es_administrador:
                    sql = f"""
                        SELECT b.rut, b.dj_numero, b.periodo, b.observacion_code, b.glosa_observacion, b.descripcion, b.orientacion
                        FROM {self.base_datos}.dj_integral_observaciones b
                        INNER JOIN {self.base_datos}.empresas e ON e.run_rut = b.rut
                        WHERE e.auditor = %s
                        ORDER BY b.rut, b.dj_numero, b.periodo, b.observacion_code
                    """
                    cursor.execute(sql, (nombre_usuario,))
                else:
                    sql = f"""
                        SELECT b.rut, b.dj_numero, b.periodo, b.observacion_code, b.glosa_observacion, b.descripcion, b.orientacion
                        FROM {self.base_datos}.dj_integral_observaciones b
                        ORDER BY b.rut, b.dj_numero, b.periodo, b.observacion_code
                    """
                    cursor.execute(sql)

                filas = cursor.fetchall()

            if not filas:
                return None


            # Importes para openpyxl y typing (locales para evitar import global innecesario)
            from openpyxl import Workbook
            from openpyxl.utils import get_column_letter
            from openpyxl.worksheet.worksheet import Worksheet
            from typing import cast

            wb = Workbook()
            # Forzar tipo para que Pylance reconozca métodos y atributos
            ws = cast(Worksheet, wb.active)
            ws.title = 'Observaciones DJ'

            # Encabezados (agregamos Glosa_Observacion después de Observacion_Code)
            headers = ['RUT', 'DJ_Numero', 'Periodo', 'Observacion', 'Glosa_Observacion', 'Descripcion', 'Orientacion']
            ws.append(headers)

            for fila in filas:
                ws.append([
                    fila.get('rut'),
                    fila.get('dj_numero'),
                    fila.get('periodo'),
                    fila.get('observacion_code'),
                    fila.get('glosa_observacion'),
                    fila.get('descripcion'),
                    fila.get('orientacion')
                ])

            # Autoajustar anchos (simple) usando get_column_letter para evitar problemas con MergedCell
            for idx, column_cells in enumerate(ws.columns, start=1):
                # Calcular largo máximo en la columna
                length = 0
                for cell in column_cells:
                    try:
                        val = cell.value
                        l = len(str(val)) if val is not None else 0
                    except Exception:
                        l = 0
                    if l > length:
                        length = l

                adjusted_width = length + 2
                col_letter = get_column_letter(idx)
                ws.column_dimensions[col_letter].width = adjusted_width

            buffer = BytesIO()
            wb.save(buffer)
            buffer.seek(0)
            return buffer

        except Exception as e:
            print(f"ERROR: exportar_observaciones_dj_excel: {e}")
            import traceback
            traceback.print_exc()
            return None
