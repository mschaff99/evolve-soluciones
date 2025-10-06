"""
Servicio de Consulta Integral para Evolve Soluciones
====================================================

Contiene toda la lógica de negocio para la funcionalidad de consulta integral F29,
incluyendo obtención de datos, filtros y exportación.
"""

import pymysql
from datetime import datetime
import re
from io import BytesIO
import openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side

from aplicacion.modelos.base_datos import obtener_conexion_local, obtener_conexion_remota
from configuracion.configuracion import obtener_configuracion

# Obtener configuración
config = obtener_configuracion()


class ServicioConsultaIntegral:
    """Servicio para manejar consultas integrales F29"""

    def __init__(self):
        self.config = config

    def obtener_conexion_evolve(self):
        """
        Obtiene conexión a la base de datos evolve (local)

        Returns:
            pymysql.Connection: Conexión a la base de datos
        """
        return obtener_conexion_local()

    def limpiar_rut(self, rut):
        """
        Limpia el RUT removiendo puntos y guiones

        Args:
            rut (str): RUT con formato

        Returns:
            str: RUT limpio sin puntos ni guiones
        """
        if not rut:
            return ""
        return re.sub(r'[.\-]', '', str(rut))

    def obtener_datos_empresas_con_periodos(self, filtros=None):
        """
        Obtiene datos de empresas con sus períodos tributarios

        Args:
            filtros (dict): Filtros a aplicar en la consulta

        Returns:
            list: Lista de empresas con sus períodos
        """
        if filtros is None:
            filtros = {}

        conexion = None
        try:
            conexion = self.obtener_conexion_evolve()

            # Construir consulta base
            consulta_base = """
                SELECT DISTINCT
                    e.run_rut as rut,
                    e.empresa as nombre,
                    e.auditor as usuario,
                    e.grupo as grupo,
                    ci.id as consulta_id,
                    ci.periodo,
                    ci.tabla_resultados,
                    ci.estado,
                    COALESCE(obs_count.total_observaciones, 0) as total_observaciones,
                    GROUP_CONCAT(DISTINCT obs_codigos.codigo ORDER BY obs_codigos.codigo SEPARATOR ',') as codigos_observaciones,
                    GROUP_CONCAT(DISTINCT CONCAT(obs_codigos.codigo, ':', COALESCE(obs_codigos.descripcion, '')) ORDER BY obs_codigos.codigo SEPARATOR '|') as observaciones_detalle
                FROM empresas e
                LEFT JOIN stratex.consulta_integral ci ON e.run_rut = ci.rut
                LEFT JOIN (
                    SELECT consulta_id, COUNT(*) as total_observaciones
                    FROM stratex.observaciones
                    GROUP BY consulta_id
                ) obs_count ON ci.id = obs_count.consulta_id
                LEFT JOIN stratex.observaciones obs_codigos ON ci.id = obs_codigos.consulta_id
                WHERE 1=1
            """

            parametros = []

            # Aplicar filtros
            if filtros.get('empresa_filtro'):
                consulta_base += " AND e.empresa LIKE %s"
                parametros.append(f"%{filtros['empresa_filtro']}%")

            if filtros.get('rut_filtro'):
                rut_limpio = self.limpiar_rut(filtros['rut_filtro'])
                consulta_base += " AND (e.run_rut LIKE %s OR e.run_rut LIKE %s)"
                parametros.extend([f"%{filtros['rut_filtro']}%", f"%{rut_limpio}%"])

            if filtros.get('usuario_filtro'):
                consulta_base += " AND e.auditor = %s"
                parametros.append(filtros['usuario_filtro'])

            if filtros.get('estado_filtro'):
                consulta_base += " AND ci.estado = %s"
                parametros.append(filtros['estado_filtro'])

            # Filtros de año
            if filtros.get('año_desde'):
                consulta_base += " AND YEAR(STR_TO_DATE(CONCAT(ci.periodo, '01'), '%Y%m%d')) >= %s"
                parametros.append(filtros['año_desde'])

            if filtros.get('año_hasta'):
                consulta_base += " AND YEAR(STR_TO_DATE(CONCAT(ci.periodo, '01'), '%Y%m%d')) <= %s"
                parametros.append(filtros['año_hasta'])

            consulta_base += " GROUP BY e.run_rut, e.empresa, e.auditor, e.grupo, ci.periodo, ci.tabla_resultados, ci.estado, obs_count.total_observaciones"
            consulta_base += " ORDER BY e.empresa, ci.periodo"

            with conexion.cursor(pymysql.cursors.DictCursor) as cursor:
                cursor.execute(consulta_base, parametros)
                resultados = cursor.fetchall()

            # Agrupar resultados por empresa
            empresas_agrupadas = {}

            for fila in resultados:
                rut = fila['rut']

                if rut not in empresas_agrupadas:
                    empresas_agrupadas[rut] = {
                        'rut': rut,
                        'nombre': fila['nombre'],
                        'usuario': fila['usuario'],
                        'grupo': fila['grupo'],
                        'periodos': {}
                    }

                # Agregar período si existe
                if fila['periodo']:
                    periodo_str = str(fila['periodo'])
                    if len(periodo_str) == 6:  # YYYYMM
                        año = int(periodo_str[:4])
                        mes = int(periodo_str[4:])

                        empresas_agrupadas[rut]['periodos'][(año, mes)] = {
                            'periodo': fila['periodo'],
                            'tabla_resultados': fila['tabla_resultados'] or '',
                            'estado': fila['estado'] or '',
                            'total_observaciones': fila['total_observaciones'] or 0,
                            'codigos_observaciones': fila.get('codigos_observaciones', '').split(',') if fila.get('codigos_observaciones') else [],
                            'observaciones_detalle': fila.get('observaciones_detalle', '')
                        }

            return list(empresas_agrupadas.values())

        except Exception as e:
            print(f"Error obteniendo datos de empresas: {e}")
            raise
        finally:
            if conexion:
                conexion.close()

    def obtener_detalles_empresa_por_rut(self, rut):
        """
        Obtiene detalles específicos de una empresa por su RUT

        Args:
            rut (str): RUT de la empresa

        Returns:
            dict: Detalles de la empresa
        """
        conexion = None
        try:
            conexion = self.obtener_conexion_evolve()

            consulta = """
                SELECT
                    run_rut as rut,
                    empresa as nombre,
                    quien_registra as usuario,
                    grupo,
                    fecha_registro,
                    activo
                FROM empresas
                WHERE run_rut = %s
            """

            with conexion.cursor(pymysql.cursors.DictCursor) as cursor:
                cursor.execute(consulta, (rut,))
                resultado = cursor.fetchone()

            if resultado:
                return {
                    'rut': resultado['rut'],
                    'nombre': resultado['nombre'],
                    'usuario': resultado['usuario'],
                    'grupo': resultado['grupo'],
                    'fecha_registro': resultado['fecha_registro'].isoformat() if resultado['fecha_registro'] else None,
                    'activo': resultado['activo']
                }
            else:
                return None

        except Exception as e:
            print(f"Error obteniendo detalles de empresa {rut}: {e}")
            raise
        finally:
            if conexion:
                conexion.close()

    def obtener_estadisticas_generales(self):
        """
        Obtiene estadísticas generales del sistema

        Returns:
            dict: Estadísticas generales
        """
        conexion = None
        try:
            conexion = self.obtener_conexion_evolve()

            with conexion.cursor(pymysql.cursors.DictCursor) as cursor:
                # Total de empresas
                cursor.execute("SELECT COUNT(*) as total FROM empresas WHERE activo = 1")
                total_empresas = cursor.fetchone()['total']  # type: ignore

                # Total de períodos
                cursor.execute("SELECT COUNT(*) as total FROM stratex.consulta_integral")
                total_periodos = cursor.fetchone()['total']  # type: ignore

                # Total de observaciones
                cursor.execute("SELECT COUNT(*) as total FROM stratex.observaciones")
                total_observaciones = cursor.fetchone()['total']  # type: ignore

                # Períodos por estado
                cursor.execute("""
                    SELECT estado, COUNT(*) as cantidad
                    FROM stratex.consulta_integral
                    GROUP BY estado
                """)
                periodos_por_estado = cursor.fetchall()

                # Empresas por usuario
                cursor.execute("""
                    SELECT quien_registra as usuario, COUNT(*) as cantidad
                    FROM empresas
                    WHERE activo = 1 AND quien_registra IS NOT NULL
                    GROUP BY quien_registra
                    ORDER BY cantidad DESC
                """)
                empresas_por_usuario = cursor.fetchall()

            return {
                'total_empresas': total_empresas,
                'total_periodos': total_periodos,
                'total_observaciones': total_observaciones,
                'periodos_por_estado': periodos_por_estado,
                'empresas_por_usuario': empresas_por_usuario,
                'fecha_generacion': datetime.now().isoformat()
            }

        except Exception as e:
            print(f"Error obteniendo estadísticas: {e}")
            raise
        finally:
            if conexion:
                conexion.close()

    def obtener_codigos_observaciones_unicos(self):
        """
        Obtiene los códigos únicos de observaciones desde el catálogo en PostgreSQL

        Returns:
            list: Lista de diccionarios con código y descripción
        """
        try:
            # Importar función de conexión a PostgreSQL
            from aplicacion.modelos.base_datos import ejecutar_consulta_postgres

            consulta = """
                SELECT
                    codigo,
                    descripcion
                FROM codigos_observaciones_f29
                WHERE activo = TRUE
                ORDER BY codigo
            """

            # Ejecutar consulta en PostgreSQL
            resultados = ejecutar_consulta_postgres(consulta)

            # Validar que se obtuvieron resultados
            if not resultados:
                print("[WARN] No se encontraron códigos en el catálogo PostgreSQL")
                return []

            # Retornar lista completa con toda la información
            codigos = []
            for fila in resultados:
                codigos.append({
                    'codigo': fila['codigo'],
                    'descripcion': fila['descripcion']
                })

            print(f"[OK] Códigos de observaciones encontrados en catálogo PostgreSQL: {len(codigos)}")
            if codigos:
                print(f"     Primeros códigos: {', '.join([c['codigo'] for c in codigos[:5]])}{'...' if len(codigos) > 5 else ''}")

            return codigos

        except Exception as e:
            print(f"[ERROR] Error obteniendo códigos de observaciones desde catálogo PostgreSQL: {e}")
            import traceback
            traceback.print_exc()

            # Fallback: intentar obtener desde observaciones existentes en MySQL si falla PostgreSQL
            print("        Intentando fallback desde tabla observaciones (MySQL)...")
            conexion_fallback = None
            try:
                conexion_fallback = self.obtener_conexion_evolve()
                consulta_fallback = """
                    SELECT DISTINCT codigo
                    FROM observaciones
                    WHERE codigo IS NOT NULL AND codigo != ''
                    ORDER BY codigo
                """
                with conexion_fallback.cursor(pymysql.cursors.DictCursor) as cursor:
                    cursor.execute(consulta_fallback)
                    resultados_fallback = cursor.fetchall()

                # Retornar en formato simple para compatibilidad
                codigos_fallback = [{'codigo': fila['codigo'], 'descripcion': ''}
                                   for fila in resultados_fallback]

                print(f"        [OK] Fallback exitoso: {len(codigos_fallback)} códigos desde observaciones MySQL")
                return codigos_fallback

            except Exception as e_fallback:
                print(f"        [ERROR] Error en fallback: {e_fallback}")
                # Retornar lista vacía como último recurso
                return []
            finally:
                if conexion_fallback:
                    conexion_fallback.close()

    def exportar_observaciones_usuario_excel(self, nombre_usuario, es_administrador=False):
        """
        Exporta las observaciones de un usuario específico a Excel
        Si es administrador, exporta TODAS las observaciones

        Args:
            nombre_usuario (str): Nombre del usuario (auditor)
            es_administrador (bool): Si es True, exporta todas las observaciones

        Returns:
            BytesIO: Buffer con el archivo Excel
        """
        conexion = None
        try:
            conexion = self.obtener_conexion_evolve()

            # Consulta para obtener observaciones
            # Si es administrador, no filtra por auditor
            if es_administrador:
                consulta = """
                    SELECT
                        c.empresa as nombre_empresa,
                        c.run_rut as rut,
                        b.periodo,
                        b.año,
                        b.mes,
                        a.codigo,
                        a.descripcion,
                        a.monto,
                        c.auditor
                    FROM observaciones a
                    LEFT JOIN consulta_integral b ON a.consulta_id = b.id
                    LEFT JOIN empresas c ON b.rut = c.run_rut
                    ORDER BY c.empresa, b.periodo, a.codigo
                """
                parametros = ()
            else:
                # Usuario normal: solo sus observaciones
                consulta = """
                    SELECT
                        c.empresa as nombre_empresa,
                        c.run_rut as rut,
                        b.periodo,
                        b.año,
                        b.mes,
                        a.codigo,
                        a.descripcion,
                        a.monto
                    FROM observaciones a
                    LEFT JOIN consulta_integral b ON a.consulta_id = b.id
                    LEFT JOIN empresas c ON b.rut = c.run_rut
                    WHERE c.auditor = %s
                    ORDER BY c.empresa, b.periodo, a.codigo
                """
                parametros = (nombre_usuario,)

            with conexion.cursor(pymysql.cursors.DictCursor) as cursor:
                cursor.execute(consulta, parametros)
                resultados = cursor.fetchall()

            if not resultados:
                return None

            # Crear workbook
            wb = openpyxl.Workbook()
            ws = wb.active  # type: ignore

            # Título según tipo de usuario
            if es_administrador:
                ws.title = "Todas las Observaciones"  # type: ignore
            else:
                ws.title = f"Observaciones {nombre_usuario}"  # type: ignore

            # Estilos
            header_font = Font(bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="1F4788", end_color="1F4788", fill_type="solid")
            border = Border(
                left=Side(style='thin'),
                right=Side(style='thin'),
                top=Side(style='thin'),
                bottom=Side(style='thin')
            )

            # Headers (agregar columna Auditor si es administrador)
            if es_administrador:
                headers = ['Empresa', 'RUT', 'Período', 'Año', 'Mes', 'Código', 'Descripción', 'Monto', 'Auditor']
            else:
                headers = ['Empresa', 'RUT', 'Período', 'Año', 'Mes', 'Código', 'Descripción', 'Monto']

            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col, value=header)  # type: ignore
                cell.font = header_font
                cell.fill = header_fill
                cell.border = border

            # Datos
            for row_idx, fila in enumerate(resultados, start=2):
                ws.cell(row=row_idx, column=1, value=fila['nombre_empresa']).border = border  # type: ignore
                ws.cell(row=row_idx, column=2, value=fila['rut']).border = border  # type: ignore
                ws.cell(row=row_idx, column=3, value=fila['periodo']).border = border  # type: ignore
                ws.cell(row=row_idx, column=4, value=fila['año']).border = border  # type: ignore

                # Mes ya viene como nombre desde la BD (Enero, Febrero, etc.)
                ws.cell(row=row_idx, column=5, value=fila.get('mes', '')).border = border  # type: ignore

                ws.cell(row=row_idx, column=6, value=fila['codigo']).border = border  # type: ignore
                ws.cell(row=row_idx, column=7, value=fila['descripcion']).border = border  # type: ignore
                ws.cell(row=row_idx, column=8, value=fila['monto']).border = border  # type: ignore

                # Agregar columna Auditor si es administrador
                if es_administrador:
                    ws.cell(row=row_idx, column=9, value=fila.get('auditor', '')).border = border  # type: ignore

            # Ajustar ancho de columnas
            if es_administrador:
                column_widths = {
                    1: 30,  # Empresa
                    2: 15,  # RUT
                    3: 12,  # Período
                    4: 10,  # Año
                    5: 12,  # Mes
                    6: 12,  # Código
                    7: 50,  # Descripción
                    8: 15,  # Monto
                    9: 20   # Auditor
                }
            else:
                column_widths = {
                    1: 30,  # Empresa
                    2: 15,  # RUT
                    3: 12,  # Período
                    4: 10,  # Año
                    5: 12,  # Mes
                    6: 12,  # Código
                    7: 50,  # Descripción
                    8: 15   # Monto
                }

            for col, width in column_widths.items():
                ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = width  # type: ignore

            # Guardar en buffer
            buffer = BytesIO()
            wb.save(buffer)
            buffer.seek(0)

            return buffer

        except Exception as e:
            print(f"Error exportando observaciones a Excel: {e}")
            import traceback
            traceback.print_exc()
            return None
        finally:
            if conexion:
                conexion.close()

    def exportar_a_excel(self, datos_empresas):
        """
        Exporta los datos de empresas a un archivo Excel

        Args:
            datos_empresas (list): Lista de empresas con sus datos

        Returns:
            BytesIO: Buffer con el archivo Excel generado
        """
        try:
            # Crear workbook
            wb = openpyxl.Workbook()
            ws = wb.active
            if ws:  # type: ignore
                ws.title = "Consulta Integral F29"  # type: ignore

            # Estilos
            header_font = Font(bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="B91C1C", end_color="B91C1C", fill_type="solid")
            border = Border(
                left=Side(style='thin'),
                right=Side(style='thin'),
                top=Side(style='thin'),
                bottom=Side(style='thin')
            )

            # Headers
            headers = [
                'RUT', 'Empresa', 'Usuario', 'Grupo', 'Año', 'Mes',
                'Período', 'Resultado', 'Estado', 'Observaciones'
            ]

            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col, value=header)  # type: ignore
                cell.font = header_font
                cell.fill = header_fill
                cell.border = border

            # Datos
            row = 2
            meses_nombres = [
                'Enero', 'Febrero', 'Marzo', 'Abril', 'Mayo', 'Junio',
                'Julio', 'Agosto', 'Septiembre', 'Octubre', 'Noviembre', 'Diciembre'
            ]

            for empresa in datos_empresas:
                for (año, mes), periodo_data in empresa['periodos'].items():
                    ws.cell(row=row, column=1, value=empresa['rut']).border = border  # type: ignore
                    ws.cell(row=row, column=2, value=empresa['nombre']).border = border  # type: ignore
                    ws.cell(row=row, column=3, value=empresa['usuario']).border = border  # type: ignore
                    ws.cell(row=row, column=4, value=empresa['grupo']).border = border  # type: ignore
                    ws.cell(row=row, column=5, value=año).border = border  # type: ignore
                    ws.cell(row=row, column=6, value=meses_nombres[mes-1]).border = border  # type: ignore
                    ws.cell(row=row, column=7, value=periodo_data['periodo']).border = border  # type: ignore
                    ws.cell(row=row, column=8, value=periodo_data['tabla_resultados']).border = border  # type: ignore
                    ws.cell(row=row, column=9, value=periodo_data['estado']).border = border  # type: ignore
                    ws.cell(row=row, column=10, value=periodo_data['total_observaciones']).border = border  # type: ignore
                    row += 1

            # Ajustar ancho de columnas
            for col in range(1, len(headers) + 1):
                ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 15  # type: ignore

            # Guardar en buffer
            buffer = BytesIO()
            wb.save(buffer)
            buffer.seek(0)

            return buffer

        except Exception as e:
            print(f"Error exportando a Excel: {e}")
            return None
